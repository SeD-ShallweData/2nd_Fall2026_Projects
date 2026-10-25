"""SK 제공 점의 위치 인덱스 / 고정 기준일 경계 매핑 / 원본에 결과 붙이기.

원본을 읽기 전용으로 사용한다. index에는 CRS/경계가 필요하지 않다.
실행 예시는 프로젝트 루트의 README.md 참고. Python 3.11 이상.
"""
from __future__ import annotations

import argparse
import codecs
import csv
import hashlib
import itertools
import json
import math
import random
import re
import sqlite3
import sys
import zipfile
from collections import Counter, defaultdict
from contextlib import contextmanager
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path
from uuid import uuid4

import pandas as pd

ORIGINAL = ["STD_YM", "BLOCK_CD", "X_COORD", "Y_COORD"]
KEYS = ["STD_YM_KEY", "BLOCK_CD_KEY", "X_COORD_KEY", "Y_COORD_KEY"]
ROLES = ["month", "block", "x", "y"]
NORMALIZATION_VERSION = "decimal-exact-v1"
INVALID = "!INVALID:"


def dump_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def csv_writer(path, fields):
    stream = Path(path).open("w", encoding="utf-8-sig", newline="")
    writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="raise")
    writer.writeheader()
    return stream, writer


def sql_csv(db, query, path, parameters=()):
    cursor = db.execute(query, parameters)
    with Path(path).open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow([d[0] for d in cursor.description])
        writer.writerows(cursor)


def run_dir(config, kind):
    # 새 기준 노트북은 사전에 확인·정리한 고정 결과 폴더를 사용한다.
    if config.get("_run_directory"):
        out = Path(config["_run_directory"]).resolve()
        allowed = (Path(config["output_dir"]) / "canonical").resolve()
        if not out.is_relative_to(allowed) or out == allowed:
            raise ValueError("고정 실행 폴더는 canonical 하위 경로여야 합니다.")
        out.mkdir(parents=True, exist_ok=False)
        dump_json(out / "config_used.json", config)
        return out
    root = Path(config["output_dir"]).resolve()
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    out = root / f"{kind}_{stamp}_{uuid4().hex[:8]}"
    out.mkdir(parents=True, exist_ok=False)
    dump_json(out / "config_used.json", config)
    return out


def readonly_db(path):
    return sqlite3.connect(Path(path).resolve().as_uri() + "?mode=ro", uri=True)


def month_key(value):
    text = str(value).strip()
    match = re.fullmatch(r"(\d{4})[-/.]?(\d{2})", text)
    if match and 1 <= int(match[2]) <= 12 and 1 <= int(match[1]) <= 9999:
        return match[1] + match[2]
    return INVALID + text


def decimal_key(value):
    """Decimal.normalize()의 기본 정밀도 반올림을 피하고 숫자 의미를 보존한다."""
    text = str(value).strip()
    if not text or len(text) > 2048:
        return INVALID + text
    try:
        number = Decimal(text)
    except InvalidOperation:
        return INVALID + text
    if not number.is_finite():
        return INVALID + text
    if number == 0:
        return "0"
    sign, digits, exponent = number.as_tuple()
    digits = "".join(str(digit) for digit in digits).lstrip("0")
    trimmed = digits.rstrip("0")
    exponent += len(digits) - len(trimmed)
    return ("-" if sign else "") + trimmed + "e" + str(exponent)


def normalize_values(raw):
    return (month_key(raw[0]), str(raw[1]).strip(), decimal_key(raw[2]), decimal_key(raw[3]))


def finite_float(key):
    if key.startswith(INVALID):
        return None
    try:
        result = float(Decimal(key))
    except (InvalidOperation, OverflowError, ValueError):
        return None
    if not math.isfinite(result) or (result == 0 and Decimal(key) != 0):
        return None
    return result


@contextmanager
def open_binary(source):
    if source.get("member"):
        with zipfile.ZipFile(source["path"]) as archive:
            with archive.open(source["member"], "r") as stream:
                yield stream
    else:
        with Path(source["path"]).open("rb") as stream:
            yield stream


def detect_csv(source, config, required=None):
    with open_binary(source) as stream:
        prefix = stream.read(131072)
    encodings = [config["input_encoding"]] if config.get("input_encoding") else config.get("encoding_candidates", ["utf-8-sig", "cp949"])
    errors = []
    for encoding in encodings:
        try:
            # prefix 끝이 멀티바이트 문자 중간이어도 잘린 부분만 다음 읽기로 넘긴다.
            text = codecs.getincrementaldecoder(encoding)(errors="strict").decode(prefix, final=False)
            separators = [config["input_separator"]] if config.get("input_separator") else [",", "|", "\t", ";"]
            seen_headers = []
            for sep in separators:
                if not isinstance(sep, str) or len(sep) != 1:
                    raise ValueError("input_separator는 한 글자여야 합니다. 탭은 JSON의 \\t로 지정합니다.")
                header = next(csv.reader(text.splitlines(keepends=True), delimiter=sep))
                seen_headers.append(header)
                names = [name.lstrip("\ufeff").strip().upper() for name in header]
                if len(header) < 2 or (required and not set(name.upper() for name in required).issubset(names)):
                    continue
                if len(names) != len(set(names)):
                    raise ValueError("BOM/공백/대소문자 정리 후 열 이름이 중복됩니다.")
                return encoding, sep, header
            raise ValueError(f"헤더에서 구분자/필수 열을 찾지 못했습니다: expected={required}, candidate_headers={seen_headers}")
        except (UnicodeError, csv.Error, StopIteration, ValueError, LookupError) as error:
            errors.append(f"{encoding}: {error}")
    raise ValueError("인코딩/구분자/필수 열을 읽지 못했습니다. 설정을 채워주세요. " + " | ".join(errors))


def resolve_columns(header, requested):
    lookup = {name.lstrip("\ufeff").strip().upper(): name for name in header}
    missing = [name for name in requested if name.upper() not in lookup]
    if missing:
        raise ValueError(f"설정한 열이 없습니다: {missing}; 실제 열: {header}")
    return [lookup[name.upper()] for name in requested]


def table_for(name, config):
    for table in config["tables"]:
        match = re.fullmatch(table["filename_regex"], name, re.IGNORECASE)
        if match:
            return table, match[1] if match.lastindex else ""
    return None, ""


def discover(config, out):
    root = Path(config["input_dir"]).resolve()
    if not root.is_dir():
        raise FileNotFoundError(root)
    sources, inventory = [], []
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        inventory.append({"path": str(path), "member": "", "bytes": path.stat().st_size, "extension": path.suffix.lower()})
        if path.suffix.lower() == ".zip":
            with zipfile.ZipFile(path) as archive:
                for item in archive.infolist():
                    if item.is_dir():
                        continue
                    inventory.append({"path": str(path), "member": item.filename, "bytes": item.file_size, "extension": Path(item.filename).suffix.lower()})
                    table, filename_month = table_for(Path(item.filename).name, config)
                    if table:
                        sources.append({"path": str(path), "member": item.filename, "table": table, "filename_month": filename_month})
        else:
            table, filename_month = table_for(path.name, config)
            if table:
                sources.append({"path": str(path), "member": "", "table": table, "filename_month": filename_month})
    stream, writer = csv_writer(out / "source_inventory.csv", ["path", "member", "bytes", "extension"])
    with stream:
        writer.writerows(inventory)
    if not sources:
        raise ValueError("tables.filename_regex에 일치하는 CSV가 없습니다. 설정을 확인해주세요.")
    return sources


def export_layout(config, out):
    """정의서의 셀과 주석을 기록한다. 기본키나 좌표계를 자동 확정하지 않는다."""
    path = Path(config["layout_file"]) if config.get("layout_file") else None
    if path is None or not path.exists():
        dump_json(out / "layout_status.json", {"status": "not_available", "path": str(path)})
        return
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=False, data_only=False)
    stream, writer = csv_writer(out / "layout_cells.csv", ["file", "sheet", "cell", "value", "comment"])
    try:
        with stream:
            for sheet in workbook:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None or cell.comment:
                            writer.writerow({"file": str(path), "sheet": sheet.title, "cell": cell.coordinate, "value": cell.value, "comment": cell.comment.text if cell.comment else ""})
    finally:
        workbook.close()


def init_index(db):
    db.executescript("""
        PRAGMA journal_mode=WAL;
        CREATE TABLE observations (
            file_id INTEGER, table_name TEXT, ym TEXT, block TEXT, x TEXT, y TEXT,
            raw_ym TEXT, raw_block TEXT, raw_x TEXT, raw_y TEXT,
            n_rows INTEGER, first_row INTEGER,
            PRIMARY KEY(file_id,ym,block,x,y,raw_ym,raw_block,raw_x,raw_y)
        );
        CREATE TABLE pk_counts (file_id INTEGER,key_json TEXT,n_rows INTEGER,first_row INTEGER,
            PRIMARY KEY(file_id,key_json));
    """)


def inspect_source(db, source, file_id, config, issues):
    table = source["table"]
    requested = [config["columns"][role] for role in ROLES]
    pk_requested = table.get("primary_key", [])
    encoding, sep, header = detect_csv(source, config, requested + pk_requested)
    columns = resolve_columns(header, requested)
    pk_columns = resolve_columns(header, pk_requested)
    used = list(dict.fromkeys(columns + pk_columns))
    report = {"file_id": file_id, "source": source["path"], "zip_member": source["member"], "table": table["name"],
              "filename_month": source["filename_month"], "encoding": encoding, "separator": repr(sep),
              "columns": json.dumps(header, ensure_ascii=False), "read_columns": json.dumps(used, ensure_ascii=False),
              "configured_pk": json.dumps(pk_columns, ensure_ascii=False), "row_meaning": table.get("row_meaning"),
              "rows_read": 0, "missing_key_rows": 0, "invalid_month_rows": 0, "invalid_coordinate_rows": 0,
              "filename_month_mismatch_rows": 0, "block_length_min": None, "block_length_max": None, "leading_zero_block_rows": 0,
              "month_min": None, "month_max": None, "months": "", "status": "complete", "error": ""}
    months, block_lengths = set(), set()
    role_by_actual = dict(zip(columns, ROLES))
    row_number = 0
    try:
        with open_binary(source) as stream:
            reader = pd.read_csv(stream, encoding=encoding, sep=sep, dtype=str, usecols=used,
                                 keep_default_na=False, na_filter=False, on_bad_lines="error", chunksize=config["chunksize"])
            for frame in reader:
                counts, first_rows, pk_counts, pk_first = Counter(), {}, Counter(), {}
                for values in frame[columns].itertuples(index=False, name=None):
                    row_number += 1
                    raw = tuple(str(value) for value in values)
                    keys = normalize_values(raw)
                    record = (file_id, table["name"], *keys, *raw)
                    counts[record] += 1
                    first_rows.setdefault(record, row_number)
                    problems = []
                    if any(not value.strip() for value in raw):
                        report["missing_key_rows"] += 1
                        problems.append("MISSING_LOCATION_KEY")
                    if keys[0].startswith(INVALID):
                        report["invalid_month_rows"] += 1
                        problems.append("INVALID_MONTH")
                    else:
                        months.add(keys[0])
                        if source["filename_month"] and keys[0] != source["filename_month"]:
                            report["filename_month_mismatch_rows"] += 1
                            problems.append("FILENAME_MONTH_MISMATCH")
                    if finite_float(keys[2]) is None or finite_float(keys[3]) is None:
                        report["invalid_coordinate_rows"] += 1
                        problems.append("INVALID_OR_NONFINITE_FLOAT_COORDINATE")
                    if keys[1]:
                        block_lengths.add(len(keys[1]))
                        report["leading_zero_block_rows"] += int(keys[1].startswith("0"))
                    if problems:
                        issues.writerow({"file_id": file_id, "source": source["path"], "member": source["member"], "row_number": row_number,
                                         "issue": "|".join(problems), **dict(zip(ORIGINAL, raw)), "detail": ""})
                if pk_columns:
                    start = row_number - len(frame) + 1
                    for offset, values in enumerate(frame[pk_columns].itertuples(index=False, name=None)):
                        normalized = []
                        for column, value in zip(pk_columns, values):
                            role = role_by_actual.get(column)
                            normalized.append(month_key(value) if role == "month" else decimal_key(value) if role in ("x", "y") else str(value).strip())
                        pk = json.dumps(normalized, ensure_ascii=False, separators=(",", ":"))
                        pk_counts[pk] += 1
                        pk_first.setdefault(pk, start + offset)
                        if any(not str(value).strip() for value in values):
                            issues.writerow({"file_id": file_id, "source": source["path"], "member": source["member"], "row_number": start + offset,
                                             "issue": "MISSING_CONFIGURED_PK", **dict(zip(ORIGINAL, [str(frame.iloc[offset][c]) for c in columns])), "detail": pk})
                db.executemany("""INSERT INTO observations VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
                    ON CONFLICT(file_id,ym,block,x,y,raw_ym,raw_block,raw_x,raw_y)
                    DO UPDATE SET n_rows=n_rows+excluded.n_rows""", [(*key, count, first_rows[key]) for key, count in counts.items()])
                db.executemany("""INSERT INTO pk_counts VALUES (?,?,?,?)
                    ON CONFLICT(file_id,key_json) DO UPDATE SET n_rows=n_rows+excluded.n_rows""",
                    [(file_id, key, count, pk_first[key]) for key, count in pk_counts.items()])
                db.commit()
                print(f"  {Path(source['path']).name}/{source['member']} : {row_number:,}행", flush=True)
    except Exception as error:
        report["status"], report["error"] = "failed_partial", str(error)
    report["rows_read"] = row_number
    report["months"] = "|".join(sorted(months))
    report["month_min"], report["month_max"] = (min(months), max(months)) if months else (None, None)
    report["block_length_min"], report["block_length_max"] = (min(block_lengths), max(block_lengths)) if block_lengths else (None, None)
    return report


def finalize_index(db, out, table_names):
    db.executescript("""
        CREATE INDEX observation_location ON observations(ym,block,x,y);
        CREATE TABLE locations (
            location_id INTEGER PRIMARY KEY,ym TEXT,block TEXT,x TEXT,y TEXT,
            raw_ym TEXT,raw_block TEXT,raw_x TEXT,raw_y TEXT,
            source_rows INTEGER,file_count INTEGER,table_count INTEGER,
            UNIQUE(ym,block,x,y)
        );
        INSERT INTO locations(ym,block,x,y,raw_ym,raw_block,raw_x,raw_y,source_rows,file_count,table_count)
        WITH grouped AS (
            SELECT ym,block,x,y,MIN(rowid) example_id,SUM(n_rows) total,
                   COUNT(DISTINCT file_id) nf,COUNT(DISTINCT table_name) nt
            FROM observations GROUP BY ym,block,x,y
        )
        SELECT g.ym,g.block,g.x,g.y,o.raw_ym,o.raw_block,o.raw_x,o.raw_y,g.total,g.nf,g.nt
        FROM grouped g JOIN observations o ON o.rowid=g.example_id ORDER BY g.ym,g.block,g.x,g.y;
        CREATE TABLE table_locations (
            table_name TEXT,ym TEXT,block TEXT,x TEXT,y TEXT,
            PRIMARY KEY(table_name,ym,block,x,y)
        );
        INSERT INTO table_locations SELECT DISTINCT table_name,ym,block,x,y FROM observations;
        INSERT INTO table_locations SELECT '__ALL__',ym,block,x,y FROM locations;
        CREATE INDEX table_location_month ON table_locations(table_name,ym,block);
        CREATE TABLE signatures(table_name TEXT,ym TEXT,block TEXT,n_pairs INTEGER,sha256 TEXT,
                                PRIMARY KEY(table_name,ym,block));
    """)
    cursor = db.execute("SELECT table_name,ym,block,x,y FROM table_locations ORDER BY table_name,ym,block,x,y")
    for group, rows in itertools.groupby(cursor, key=lambda row: row[:3]):
        digest, count = hashlib.sha256(), 0
        for row in rows:
            digest.update((json.dumps(row[3:5], separators=(",", ":")) + "\n").encode("utf-8"))
            count += 1
        db.execute("INSERT INTO signatures VALUES (?,?,?,?,?)", (*group, count, digest.hexdigest()))
    db.commit()
    sql_csv(db, """SELECT location_id,raw_ym STD_YM,raw_block BLOCK_CD,raw_x X_COORD,raw_y Y_COORD,
        ym STD_YM_KEY,block BLOCK_CD_KEY,x X_COORD_KEY,y Y_COORD_KEY,source_rows,file_count,table_count
        FROM locations ORDER BY location_id""", out / "location_index.csv")
    sql_csv(db, """SELECT o.file_id,o.table_name,l.location_id,o.raw_ym STD_YM,o.raw_block BLOCK_CD,
        o.raw_x X_COORD,o.raw_y Y_COORD,o.ym STD_YM_KEY,o.block BLOCK_CD_KEY,o.x X_COORD_KEY,o.y Y_COORD_KEY,o.n_rows
        FROM observations o JOIN locations l USING(ym,block,x,y) ORDER BY l.location_id,o.file_id""", out / "location_aliases.csv")
    sql_csv(db, """SELECT file_id,table_name,ym,block,x,y,SUM(n_rows) occurrences,MIN(first_row) first_row,
        'LOCATION_REPEAT_DOES_NOT_PROVE_ROW_PK_DUPLICATE' interpretation
        FROM observations GROUP BY file_id,table_name,ym,block,x,y HAVING SUM(n_rows)>1""", out / "location_key_repeats.csv")
    sql_csv(db, "SELECT * FROM pk_counts WHERE n_rows>1", out / "configured_pk_duplicates.csv")
    sql_csv(db, """SELECT f.table_name,p.key_json,SUM(p.n_rows) occurrences,COUNT(DISTINCT p.file_id) file_count,
        GROUP_CONCAT(DISTINCT p.file_id) file_ids FROM pk_counts p
        JOIN (SELECT DISTINCT file_id,table_name FROM observations) f USING(file_id)
        GROUP BY f.table_name,p.key_json HAVING SUM(p.n_rows)>1""", out / "configured_pk_duplicates_across_files.csv")
    sql_csv(db, "SELECT * FROM signatures ORDER BY table_name,ym,block", out / "monthly_block_coordinate_counts.csv")
    sql_csv(db, """SELECT file_id,table_name,ym,block,COUNT(*) coordinate_pair_count FROM
        (SELECT DISTINCT file_id,table_name,ym,block,x,y FROM observations)
        GROUP BY file_id,table_name,ym,block""", out / "file_month_block_coordinate_counts.csv")
    sql_csv(db, """SELECT table_name,ym,x,y,COUNT(*) block_count,GROUP_CONCAT(block,'|') block_codes
        FROM table_locations WHERE x NOT LIKE '!INVALID:%' AND y NOT LIKE '!INVALID:%' AND block<>''
        GROUP BY table_name,ym,x,y HAVING COUNT(*)>1""", out / "shared_coordinates_multiple_blocks.csv")
    sql_csv(db, """SELECT l.location_id,COUNT(*) raw_variants FROM
        (SELECT DISTINCT ym,block,x,y,raw_ym,raw_block,raw_x,raw_y FROM observations) o
        JOIN locations l USING(ym,block,x,y) GROUP BY l.location_id HAVING COUNT(*)>1""", out / "numeric_spelling_variants.csv")
    compare_tables(db, out, table_names)
    compare_months(db, out, [*table_names, "__ALL__"])
    db.execute("PRAGMA wal_checkpoint(TRUNCATE)")


def compare_tables(db, out, tables):
    summary_stream, summary = csv_writer(out / "table_coordinate_comparison.csv", ["left_table", "right_table", "ym", "block", "left_pairs", "right_pairs", "left_only_pairs", "right_only_pairs", "status"])
    detail_stream, detail = csv_writer(out / "table_coordinate_differences.csv", ["left_table", "right_table", "present_in", "ym", "block", "x", "y"])
    with summary_stream, detail_stream:
        for left, right in itertools.combinations(tables, 2):
            # NOT EXISTS로 실제 집합 차이를 계산한다. 해시만으로 일치를 확정하지 않는다.
            query = """WITH keys AS (
                SELECT ym,block FROM signatures WHERE table_name=? UNION SELECT ym,block FROM signatures WHERE table_name=?
            ), ld AS (
                SELECT a.ym,a.block,COUNT(*) n FROM table_locations a WHERE a.table_name=? AND NOT EXISTS
                (SELECT 1 FROM table_locations b WHERE b.table_name=? AND b.ym=a.ym AND b.block=a.block AND b.x=a.x AND b.y=a.y)
                GROUP BY a.ym,a.block
            ), rd AS (
                SELECT a.ym,a.block,COUNT(*) n FROM table_locations a WHERE a.table_name=? AND NOT EXISTS
                (SELECT 1 FROM table_locations b WHERE b.table_name=? AND b.ym=a.ym AND b.block=a.block AND b.x=a.x AND b.y=a.y)
                GROUP BY a.ym,a.block
            ) SELECT k.ym,k.block,COALESCE(l.n_pairs,0),COALESCE(r.n_pairs,0),COALESCE(ld.n,0),COALESCE(rd.n,0)
                FROM keys k LEFT JOIN signatures l ON l.table_name=? AND l.ym=k.ym AND l.block=k.block
                LEFT JOIN signatures r ON r.table_name=? AND r.ym=k.ym AND r.block=k.block
                LEFT JOIN ld ON ld.ym=k.ym AND ld.block=k.block
                LEFT JOIN rd ON rd.ym=k.ym AND rd.block=k.block ORDER BY k.ym,k.block"""
            for ym, block, nl, nr, dl, dr in db.execute(query, (left, right, left, right, right, left, left, right)):
                status = "ONLY_RIGHT_GROUP" if nl == 0 else "ONLY_LEFT_GROUP" if nr == 0 else "SAME_SET" if dl == dr == 0 else "DIFFERENT_SET_REVIEW_POPULATION_AND_STRUCTURE"
                summary.writerow(dict(zip(summary.fieldnames, (left, right, ym, block, nl, nr, dl, dr, status))))
            for present, absent in [(left, right), (right, left)]:
                for ym, block, x, y in db.execute("""SELECT a.ym,a.block,a.x,a.y FROM table_locations a WHERE a.table_name=?
                    AND NOT EXISTS (SELECT 1 FROM table_locations b WHERE b.table_name=? AND b.ym=a.ym AND b.block=a.block AND b.x=a.x AND b.y=a.y)""", (present, absent)):
                    detail.writerow(dict(zip(detail.fieldnames, (left, right, present, ym, block, x, y))))


def month_number(month):
    return int(month[:4]) * 12 + int(month[4:]) - 1


def compare_months(db, out, tables):
    fields = ["table", "previous_month", "current_month", "block", "previous_pairs", "current_pairs", "status"]
    stream, writer = csv_writer(out / "monthly_block_stability.csv", fields)
    gaps = []
    with stream:
        for table in tables:
            months = [r[0] for r in db.execute("SELECT DISTINCT ym FROM signatures WHERE table_name=? AND ym NOT LIKE '!INVALID:%' ORDER BY ym", (table,))]
            for previous, current in zip(months, months[1:]):
                if month_number(current) - month_number(previous) != 1:
                    gaps.append({"table": table, "previous_month": previous, "current_month": current, "status": "NOT_ADJACENT_CALENDAR_MONTHS"})
                    continue
                query = """WITH codes AS (SELECT block FROM signatures WHERE table_name=? AND ym=?
                    UNION SELECT block FROM signatures WHERE table_name=? AND ym=?)
                    SELECT c.block,COALESCE(p.n_pairs,0),COALESCE(n.n_pairs,0),
                        CASE WHEN p.block IS NULL THEN 'APPEARED' WHEN n.block IS NULL THEN 'DISAPPEARED'
                        WHEN EXISTS (SELECT 1 FROM table_locations a WHERE a.table_name=? AND a.ym=? AND a.block=c.block
                            AND NOT EXISTS (SELECT 1 FROM table_locations b WHERE b.table_name=? AND b.ym=? AND b.block=a.block AND b.x=a.x AND b.y=a.y))
                        OR EXISTS (SELECT 1 FROM table_locations a WHERE a.table_name=? AND a.ym=? AND a.block=c.block
                            AND NOT EXISTS (SELECT 1 FROM table_locations b WHERE b.table_name=? AND b.ym=? AND b.block=a.block AND b.x=a.x AND b.y=a.y))
                        THEN 'COORDINATE_SET_CHANGED' ELSE 'SAME_SET' END
                    FROM codes c LEFT JOIN signatures p ON p.table_name=? AND p.ym=? AND p.block=c.block
                    LEFT JOIN signatures n ON n.table_name=? AND n.ym=? AND n.block=c.block ORDER BY c.block"""
                params = (table, previous, table, current, table, previous, table, current, table, current, table, previous, table, previous, table, current)
                for block, npairs, cpairs, status in db.execute(query, params):
                    writer.writerow(dict(zip(fields, (table, previous, current, block, npairs, cpairs, status))))
    dump_json(out / "month_gaps.json", gaps)


def index_command(config):
    out = run_dir(config, "index")
    print(f"출력: {out}", flush=True)
    manifest = {"kind": "index", "status": "running", "normalization_version": NORMALIZATION_VERSION}
    dump_json(out / "manifest.json", manifest)
    sources = discover(config, out)
    export_layout(config, out)
    reports = []
    db = sqlite3.connect(out / "location_index.sqlite")
    init_index(db)
    stream, issues = csv_writer(out / "row_issues.csv", ["file_id", "source", "member", "row_number", "issue", *ORIGINAL, "detail"])
    try:
        with stream:
            for file_id, source in enumerate(sources, 1):
                try:
                    report = inspect_source(db, source, file_id, config, issues)
                except Exception as error:
                    report = {"file_id": file_id, "source": source["path"], "zip_member": source["member"], "table": source["table"]["name"], "status": "failed_before_read", "error": str(error), "rows_read": 0}
                reports.append(report)
        pd.DataFrame(reports).to_csv(out / "file_summary.csv", index=False, encoding="utf-8-sig")
        finalize_index(db, out, [table["name"] for table in config["tables"]])
        manifest.update(status="complete" if all(r["status"] == "complete" for r in reports) else "incomplete", files=len(reports),
                        rows_read=sum(r["rows_read"] for r in reports), locations=db.execute("SELECT COUNT(*) FROM locations").fetchone()[0],
                        primary_key_not_configured=[t["name"] for t in config["tables"] if not t.get("primary_key")],
                        scope="location_index_only; no CRS inference or spatial join", input_columns=config["columns"])
        dump_json(out / "manifest.json", manifest)
    finally:
        db.close()
    if manifest["status"] != "complete":
        raise ValueError(f"일부 파일 읽기가 실패했습니다. 부분 결과는 보존했습니다: {out / 'file_summary.csv'}; map은 실행되지 않습니다.")
    print(f"위치 검사 완료: {out}\n공간 변환과 공간 결합은 아직 실행하지 않았습니다.")
    return out


def require_fields(mapping, names, label):
    missing = [name for name in names if mapping.get(name) is None or mapping.get(name) == ""]
    if missing:
        raise ValueError(f"{label} 설정 필요: {', '.join(missing)}")


def parse_day(value, label):
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError as error:
        raise ValueError(f"{label}: YYYY-MM-DD 형식이 필요합니다: {value}") from error


def coordinate_settings(config):
    from pyproj import CRS
    settings = config["source_coordinates"]
    require_fields(settings, ["crs", "evidence", "unit", "xy_order", "scale_x", "scale_y"], "source_coordinates")
    crs = CRS.from_user_input(settings["crs"])
    if settings["xy_order"] not in ("xy", "yx"):
        raise ValueError("source_coordinates.xy_order는 xy 또는 yx여야 합니다.")
    for field in ("scale_x", "scale_y"):
        value = settings[field]
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value <= 0:
            raise ValueError(f"source_coordinates.{field}는 확인된 양수 배율이어야 합니다. 배율 없음이 확인되면 1.0입니다.")
    if len(crs.axis_info) < 2:
        raise ValueError("수평 2차원 CRS가 필요합니다.")
    # unit은 배율 적용 후 CRS 좌표의 단위. 데이터 배율은 scale_x/scale_y에 명시.
    factors = [axis.unit_conversion_factor for axis in crs.axis_info[:2]]
    expected = {"metre": 1.0, "degree": math.pi / 180, "foot": 0.3048, "us_survey_foot": 1200 / 3937}
    if settings["unit"] not in expected or not all(math.isclose(factor, expected[settings["unit"]], rel_tol=1e-8) for factor in factors):
        raise ValueError(f"설정한 단위와 CRS 축 단위가 다릅니다: unit={settings['unit']}, CRS axes={crs.axis_info}")
    if not (crs.is_projected or crs.is_geographic):
        raise ValueError("투영 CRS 또는 경위도 CRS를 지정해주세요.")
    return crs


def metric_settings(config, source_crs):
    from pyproj import CRS
    metric = CRS.from_user_input(config["review"]["metric_crs"]) if config["review"].get("metric_crs") else source_crs
    if not metric.is_projected or not all(math.isclose(a.unit_conversion_factor, 1.0) for a in metric.axis_info[:2]):
        raise ValueError("경계 근처 검토 샘플을 위해 review.metric_crs에 해당 지역에 적절한 미터 단위 투영 CRS를 지정해주세요.")
    return metric


def read_official_csv(settings, config, required):
    require_fields(settings, ["path", "evidence"], "공식 대응표")
    source = {"path": str(Path(settings["path"]).resolve()), "member": ""}
    overrides = {**config, "input_encoding": settings.get("encoding"), "input_separator": settings.get("separator")}
    encoding, sep, header = detect_csv(source, overrides, required)
    actual = resolve_columns(header, required)
    frame = pd.read_csv(source["path"], encoding=encoding, sep=sep, dtype=str, usecols=actual, keep_default_na=False, na_filter=False, on_bad_lines="error")
    frame = frame.rename(columns=dict(zip(actual, required)))
    # 코드에 임의의 zfill/숫자 변환은 적용하지 않는다.
    for column in required:
        frame[column] = frame[column].str.strip()
    return frame


def load_boundary(config, layer_name, out):
    import geopandas as gpd
    from pyproj import CRS
    settings = config["boundaries"][layer_name]
    require_fields(settings, ["path", "code_field", "name_field", "code_system", "reference_date", "evidence"], f"boundaries.{layer_name}")
    reference = parse_day(settings["reference_date"], f"boundaries.{layer_name}.reference_date")
    if layer_name == "emd" and settings.get("kind") not in ("administrative", "legal"):
        raise ValueError("boundaries.emd.kind에 administrative(행정동) 또는 legal(법정동)을 지정해주세요.")
    kwargs = {"engine": "pyogrio"}
    if settings.get("layer"):
        kwargs["layer"] = settings["layer"]
    if settings.get("encoding"):
        kwargs["encoding"] = settings["encoding"]
    boundary = gpd.read_file(settings["path"], **kwargs)
    code_field, name_field = settings["code_field"], settings["name_field"]
    if code_field == name_field:
        raise ValueError(f"{layer_name}: 코드 필드와 이름 필드는 서로 달라야 합니다.")
    if code_field not in boundary.columns or name_field not in boundary.columns:
        dump_json(out / f"{layer_name}_available_fields.json", list(boundary.columns))
        raise ValueError(f"{layer_name}: 설정한 코드/이름 필드가 없습니다. available_fields.json을 확인해주세요.")
    if boundary.crs is None:
        require_fields(settings, ["crs_if_missing", "crs_evidence"], f"boundaries.{layer_name}")
        boundary = boundary.set_crs(settings["crs_if_missing"])
    elif settings.get("crs_if_missing") and not CRS.from_user_input(settings["crs_if_missing"]).equals(boundary.crs):
        raise ValueError(f"{layer_name}: 파일 CRS와 crs_if_missing 설정이 다릅니다. 자동 덮어쓰지 않습니다.")
    code_is_text = boundary[code_field].map(lambda value: isinstance(value, str))
    problems = []
    for index, row in boundary.iterrows():
        geometry = row.geometry
        flags = []
        if not code_is_text.loc[index]:
            flags.append("CODE_NOT_STRING_OR_NULL")
        elif not row[code_field].strip():
            flags.append("EMPTY_CODE")
        if pd.isna(row[name_field]) or not str(row[name_field]).strip():
            flags.append("EMPTY_NAME")
        if geometry is None or geometry.is_empty:
            flags.append("EMPTY_GEOMETRY")
        elif geometry.geom_type not in ("Polygon", "MultiPolygon"):
            flags.append("NOT_POLYGON")
        elif not geometry.is_valid:
            from shapely.validation import explain_validity
            flags.append("INVALID_GEOMETRY:" + explain_validity(geometry))
        if flags:
            problems.append({"feature_index": index, "code": row[code_field], "name": row[name_field], "issue": "|".join(flags)})
    pd.DataFrame(problems, columns=["feature_index", "code", "name", "issue"]).to_csv(out / f"{layer_name}_boundary_issues.csv", index=False, encoding="utf-8-sig")
    duplicate_codes = boundary[code_field].astype(str).str.strip().duplicated(keep=False)
    boundary.loc[duplicate_codes, [code_field, name_field]].to_csv(out / f"{layer_name}_duplicate_codes.csv", index=False, encoding="utf-8-sig")
    dump_json(out / f"{layer_name}_boundary_metadata.json", {
        "path": settings["path"], "layer": settings.get("layer"), "features": len(boundary), "fields": list(boundary.columns),
        "crs_wkt": CRS.from_user_input(boundary.crs).to_wkt(), "axis_info": [str(a) for a in CRS.from_user_input(boundary.crs).axis_info],
        "reference_date": reference.isoformat(), "code_system": settings["code_system"], "kind": settings.get("kind"),
        "evidence": settings["evidence"], "duplicate_feature_codes": int(duplicate_codes.sum()), "invalid_features": len(problems)
    })
    if boundary.empty or problems or duplicate_codes.any():
        raise ValueError(f"{layer_name}: 경계가 비었거나 코드/도형 문제 또는 중복 코드가 있습니다. 진단 CSV를 확인해주세요. 자동 복구/병합은 하지 않습니다.")
    boundary = boundary[[code_field, name_field, "geometry"]].rename(columns={code_field: "region_code", name_field: "region_name"}).reset_index(drop=True)
    boundary["region_code"] = boundary["region_code"].str.strip()
    boundary["region_name"] = boundary["region_name"].astype(str).str.strip()
    return boundary


def load_parent_crosswalk(config, out):
    settings = config["parent_crosswalk"]
    if not settings.get("path"):
        dump_json(out / "parent_crosswalk_status.json", {"status": "not_provided; hierarchy cannot be confirmed"})
        return {}, set(), False
    require_fields(settings, ["emd_column", "sigungu_column", "emd_code_system", "sigungu_code_system", "reference_date"], "parent_crosswalk")
    for level in ("emd", "sigungu"):
        if settings[f"{level}_code_system"] != config["boundaries"][level]["code_system"]:
            raise ValueError(f"parent_crosswalk.{level}_code_system과 경계의 코드 체계가 다릅니다. 공식 대응표가 필요합니다.")
        if parse_day(settings["reference_date"], "parent_crosswalk.reference_date") != parse_day(config["boundaries"][level]["reference_date"], f"{level}.reference_date"):
            raise ValueError("상하위 대응표와 두 경계의 기준일이 같아야 합니다. 해당 기준일의 공식 대응표를 준비해주세요.")
    child, parent = settings["emd_column"], settings["sigungu_column"]
    frame = read_official_csv(settings, config, [child, parent])
    bad = frame[child].eq("") | frame[parent].eq("") | frame[child].duplicated(keep=False)
    frame.loc[bad].to_csv(out / "parent_crosswalk_issues.csv", index=False, encoding="utf-8-sig")
    ambiguous = set(frame.loc[bad, child])
    lookup = dict(zip(frame.loc[~bad, child], frame.loc[~bad, parent]))
    return lookup, ambiguous, True


def month_interval(month):
    start = date(int(month[:4]), int(month[4:]), 1)
    if start.month == 12:
        following = date(start.year + 1, 1, 1)
    else:
        following = date(start.year, start.month + 1, 1)
    from datetime import timedelta
    return start, following - timedelta(days=1)


def load_block_crosswalk(config, months, out):
    settings = config["block_crosswalk"]
    if not settings.get("path"):
        return {}, {}, False
    require_fields(settings, ["block_column", "sigungu_column", "emd_column", "sigungu_code_system", "emd_code_system"], "block_crosswalk")
    for level in ("sigungu", "emd"):
        if settings[f"{level}_code_system"] != config["boundaries"][level]["code_system"]:
            raise ValueError(f"block_crosswalk.{level}_code_system과 경계 코드 체계가 다릅니다.")
    required = [settings[name] for name in ("block_column", "sigungu_column", "emd_column")]
    for field in ("month_column", "valid_from_column", "valid_to_column"):
        if settings.get(field):
            required.append(settings[field])
    frame = read_official_csv(settings, config, list(dict.fromkeys(required)))
    has_period = any(settings.get(field) for field in ("valid_from_column", "valid_to_column", "valid_from", "valid_to"))
    if not settings.get("month_column") or has_period:
        if not settings.get("valid_from_column") and not settings.get("valid_from"):
            raise ValueError("공식 BLOCK 대응표의 적용 시작일을 지정해주세요.")
        if not settings.get("valid_to_column") and not settings.get("valid_to"):
            raise ValueError("공식 BLOCK 대응표의 적용 종료일을 지정해주세요.")
    candidates, partial = defaultdict(list), set()
    for record in frame.to_dict("records"):
        block = record[settings["block_column"]]
        if not block:
            raise ValueError("공식 BLOCK 대응표에 빈 BLOCK 코드가 있습니다.")
        value = (record[settings["sigungu_column"]], record[settings["emd_column"]])
        if not all(value):
            raise ValueError("공식 BLOCK 대응표의 지역 코드가 비었습니다.")
        start = end = None
        if not settings.get("month_column") or has_period:
            start = parse_day(record[settings["valid_from_column"]] if settings.get("valid_from_column") else settings["valid_from"], "block.valid_from")
            end = parse_day(record[settings["valid_to_column"]] if settings.get("valid_to_column") else settings["valid_to"], "block.valid_to")
            if start > end:
                raise ValueError("공식 BLOCK 대응표의 시작일이 종료일보다 늦습니다.")
        if settings.get("month_column"):
            month = month_key(record[settings["month_column"]])
            if month.startswith(INVALID):
                raise ValueError("공식 BLOCK 대응표의 년월 형식이 잘못되었습니다.")
            if month in months:
                first, last = month_interval(month)
                if not has_period or (start <= first and end >= last):
                    candidates[month, block].append(value)
                elif start <= last and end >= first:
                    partial.add((month, block))
        else:
            for month in months:
                first, last = month_interval(month)
                if start <= first and end >= last:
                    candidates[month, block].append(value)
                elif start <= last and end >= first:
                    partial.add((month, block))
    lookup, issues = {}, {}
    for key in set(candidates) | partial:
        if key in partial:
            issues[key] = "OFFICIAL_PARTIAL_MONTH_COVERAGE"
        elif len(candidates[key]) != 1:
            issues[key] = "OFFICIAL_BLOCK_MONTH_NOT_UNIQUE"
        else:
            lookup[key] = candidates[key][0]
    rows = [{"STD_YM_KEY": key[0], "BLOCK_CD_KEY": key[1], "issue": issue, "candidates": json.dumps(candidates.get(key, []), ensure_ascii=False)} for key, issue in issues.items()]
    pd.DataFrame(rows, columns=[*KEYS[:2], "issue", "candidates"]).to_csv(out / "block_crosswalk_issues.csv", index=False, encoding="utf-8-sig")
    return lookup, issues, True


def spatial_matches(points, boundary, layer_name, candidate_writer):
    import geopandas as gpd
    result = {int(location_id): {"count": 0, "code": "", "name": "", "contacts": []} for location_id in points.index}
    if points.empty:
        return result
    within = gpd.sjoin(points, boundary, how="inner", predicate="within")
    for location_id, matches in within.groupby(level=0):
        entries = matches[["region_code", "region_name"]].to_dict("records")
        result[int(location_id)]["count"] = len(entries)
        if len(entries) == 1:
            result[int(location_id)].update(code=entries[0]["region_code"], name=entries[0]["region_name"])
        else:
            for entry in entries:
                candidate_writer.writerow({"location_id": int(location_id), "layer": layer_name, "predicate": "within_multiple", **entry})
    unmatched = points.loc[[location_id for location_id, entry in result.items() if entry["count"] == 0]]
    if not unmatched.empty:
        contact = gpd.sjoin(unmatched, boundary, how="inner", predicate="intersects")
        for location_id, matches in contact.groupby(level=0):
            entries = matches[["region_code", "region_name"]].to_dict("records")
            result[int(location_id)]["contacts"] = entries
            for entry in entries:
                candidate_writer.writerow({"location_id": int(location_id), "layer": layer_name, "predicate": "intersects_after_no_within", **entry})
    return result


MAP_FIELDS = [
    "LOCATION_ID", "SIGUNGU_CD", "SIGUNGU_NM", "EMD_CD", "EMD_NM", "MATCH_STATUS", "REVIEW_REASONS",
    "SPATIAL_SIGUNGU_CD", "SPATIAL_SIGUNGU_NM", "SPATIAL_EMD_CD", "SPATIAL_EMD_NM",
    "SIGUNGU_WITHIN_COUNT", "EMD_WITHIN_COUNT", "SIGUNGU_CONTACT_COUNT", "EMD_CONTACT_COUNT",
    "SIGUNGU_CONTACT_CANDIDATES", "EMD_CONTACT_CANDIDATES", "HIERARCHY_STATUS", "EMD_PARENT_SIGUNGU_CD",
    "OFFICIAL_BLOCK_STATUS", "OFFICIAL_SIGUNGU_CD", "OFFICIAL_EMD_CD", "OUTSIDE_TARGET",
    "LONGITUDE", "LATITUDE", "DISTANCE_TO_BOUNDARY_M", "REVIEW_X_M", "REVIEW_Y_M",
    "SOURCE_CRS", "SOURCE_CRS_EVIDENCE", "SOURCE_UNIT", "SOURCE_XY_ORDER", "SOURCE_SCALE_X", "SOURCE_SCALE_Y",
    "POINT_MEANING", "SIGUNGU_BOUNDARY_DATE", "EMD_BOUNDARY_DATE", "SIGUNGU_BOUNDARY_CRS", "EMD_BOUNDARY_CRS",
    "SIGUNGU_CODE_SYSTEM", "EMD_CODE_SYSTEM", "EMD_KIND", "BOUNDARY_MODE", "MATCH_METHOD", "METRIC_CRS"
]


class ReviewSamples:
    """전체 점을 메모리에 모으지 않는 고정 seed 검토 표본."""
    def __init__(self, config):
        self.seed = config["review"]["seed"]
        self.random = random.Random(self.seed)
        self.n = config["review"]["samples_per_category"]
        self.near_limit = config["review"]["near_boundary_m"]
        self.targets = {target["sigungu_code"]: target["name"] for target in config["targets"]}
        self.reservoirs, self.counts, self.nearest, self.extremes = {}, Counter(), {}, {}

    def reservoir(self, bucket, record):
        self.counts[bucket] += 1
        items = self.reservoirs.setdefault(bucket, [])
        if len(items) < self.n:
            items.append(record.copy())
        else:
            index = self.random.randrange(self.counts[bucket])
            if index < self.n:
                items[index] = record.copy()

    def consider(self, record):
        code = record["SPATIAL_SIGUNGU_CD"]
        if code not in self.targets:
            return
        distance = record["DISTANCE_TO_BOUNDARY_M"]
        if distance is not None and distance > self.near_limit and record["MATCH_STATUS"] == "OK":
            self.reservoir((code, "interior"), record)
        if distance is not None:
            nearest = self.nearest.setdefault(code, [])
            nearest.append(record.copy())
            nearest.sort(key=lambda r: (r["DISTANCE_TO_BOUNDARY_M"], r["LOCATION_ID"]))
            del nearest[self.n:]
        x, y = record["REVIEW_X_M"], record["REVIEW_Y_M"]
        if x is not None and y is not None:
            for side, value, mode in [("west", x, min), ("east", x, max), ("south", y, min), ("north", y, max)]:
                key = code, side
                old = self.extremes.get(key)
                if old is None or mode(old[0], value) != old[0]:
                    self.extremes[key] = value, record.copy()
        if "춘천" in self.targets[code] and record["SPATIAL_EMD_NM"]:
            name = record["SPATIAL_EMD_NM"]
            suffix = name[-1]
            if suffix in "읍면동":
                self.reservoir((code, "chuncheon_name_suffix_" + suffix), record)

    def save(self, out):
        selection = {}
        def add(code, category, record):
            key = record["LOCATION_ID"]
            if key not in selection:
                selection[key] = {**record, "TARGET_NAME": self.targets[code], "SAMPLE_CATEGORIES": set()}
            selection[key]["SAMPLE_CATEGORIES"].add(category)
        for (code, category), items in self.reservoirs.items():
            for record in items:
                add(code, category, record)
        for code, items in self.nearest.items():
            for record in items:
                add(code, "near_boundary" if record["DISTANCE_TO_BOUNDARY_M"] <= self.near_limit else "closest_available_above_threshold", record)
        for (code, side), (_, record) in self.extremes.items():
            add(code, "distribution_extreme_" + side, record)
        fields = [*ORIGINAL, *KEYS, *MAP_FIELDS, "TARGET_NAME", "SAMPLE_CATEGORIES"]
        stream, writer = csv_writer(out / "review_samples.csv", fields)
        by_target = {target: [] for target in self.targets.values()}
        with stream:
            for record in sorted(selection.values(), key=lambda row: row["LOCATION_ID"]):
                record["SAMPLE_CATEGORIES"] = "|".join(sorted(record["SAMPLE_CATEGORIES"]))
                writer.writerow(record)
                by_target[record["TARGET_NAME"]].append(record)
        for index, (name, rows) in enumerate(by_target.items(), 1):
            path = out / f"review_samples_{index}_{name}.csv"
            stream, writer = csv_writer(path, fields)
            with stream:
                writer.writerows(rows)
        dump_json(out / "review_sampling_summary.json", {
            "seed": self.seed, "available_interior_and_suffix_counts": {"|".join(k): v for k, v in self.counts.items()},
            "selected_by_target": {name: len(rows) for name, rows in by_target.items()},
            "near_boundary_threshold_m": self.near_limit,
            "note": "표본은 월별 위치 인덱스 기준. 읍면동 구분 표본은 이름 끝 글자를 사용하며 코드 체계를 추론하지 않음. 해당 범주 점이 없으면 생성하지 않음."
        })


def find_run(config, kind, requested):
    if requested != "latest":
        result = Path(requested).resolve()
    else:
        candidates = []
        for candidate in Path(config["output_dir"]).glob(f"{kind}_*"):
            manifest_path = candidate / "manifest.json"
            if manifest_path.exists():
                metadata = json.loads(manifest_path.read_text(encoding="utf-8"))
                if metadata.get("kind") == kind and metadata.get("status") == "complete":
                    candidates.append(candidate)
        if not candidates:
            raise FileNotFoundError(f"완료된 {kind} 실행 폴더가 없습니다.")
        result = max(candidates, key=lambda path: path.name)
    metadata = json.loads((result / "manifest.json").read_text(encoding="utf-8"))
    if metadata.get("status") != "complete" or metadata.get("kind") != kind:
        raise ValueError(f"정상 완료된 {kind} 폴더가 아닙니다: {result}")
    if metadata.get("normalization_version") != NORMALIZATION_VERSION:
        raise ValueError("좌표 정규화 버전이 다릅니다. 동일 버전으로 index/map을 다시 생성해주세요.")
    return result, metadata


def point_batches(db, size):
    cursor = db.execute("SELECT * FROM locations ORDER BY location_id")
    columns = [description[0] for description in cursor.description]
    while rows := cursor.fetchmany(size):
        yield pd.DataFrame.from_records(rows, columns=columns)


def prepare_points(frame, config, transformers, boundary_crs):
    import geopandas as gpd
    from shapely.geometry import Point
    settings = config["source_coordinates"]
    ids, xs, ys = [], [], []
    for row in frame.itertuples(index=False):
        x, y = finite_float(row.x), finite_float(row.y)
        if x is None or y is None:
            continue
        x, y = x * settings["scale_x"], y * settings["scale_y"]
        if settings["xy_order"] == "yx":
            x, y = y, x
        if not math.isfinite(x) or not math.isfinite(y):
            continue
        ids.append(int(row.location_id))
        xs.append(x)
        ys.append(y)
    if not ids:
        empty = {name: gpd.GeoDataFrame(geometry=[], index=pd.Index([], dtype="int64", name="location_id"), crs=boundary_crs[name]) for name in ("sigungu", "emd")}
        return empty, {}
    transformed = {name: transformer.transform(xs, ys, errcheck=False) for name, transformer in transformers.items()}
    coordinates = {}
    geometries = {"sigungu": [], "emd": []}
    usable_ids = []
    for index, location_id in enumerate(ids):
        values = {name: (float(pair[0][index]), float(pair[1][index])) for name, pair in transformed.items()}
        longitude, latitude = values["wgs84"]
        if not all(math.isfinite(v) for pair in values.values() for v in pair) or not (-180 <= longitude <= 180 and -90 <= latitude <= 90):
            continue
        coordinates[location_id] = values
        usable_ids.append(location_id)
        for name in geometries:
            geometries[name].append(Point(*values[name]))
    points = {name: gpd.GeoDataFrame(geometry=geometries[name], index=pd.Index(usable_ids, name="location_id"), crs=boundary_crs[name]) for name in geometries}
    return points, coordinates


def mapping_attributes(row, config, source_crs, metric_crs, boundaries, matches, coordinates,
                       parent_lookup, parent_ambiguous, parent_supplied, block_lookup, block_issues, block_supplied, metric_geometry, region_names):
    from shapely.geometry import Point
    location_id = int(row.location_id)
    zero = {"count": 0, "code": "", "name": "", "contacts": []}
    sigma = matches["sigungu"].get(location_id, zero)
    emd = matches["emd"].get(location_id, zero)
    reasons = []
    attributes = {field: "" for field in MAP_FIELDS}
    attributes.update(LOCATION_ID=location_id, SPATIAL_SIGUNGU_CD=sigma["code"], SPATIAL_SIGUNGU_NM=sigma["name"],
                      SPATIAL_EMD_CD=emd["code"], SPATIAL_EMD_NM=emd["name"], SIGUNGU_WITHIN_COUNT=sigma["count"], EMD_WITHIN_COUNT=emd["count"],
                      SIGUNGU_CONTACT_COUNT=len(sigma["contacts"]), EMD_CONTACT_COUNT=len(emd["contacts"]),
                      SIGUNGU_CONTACT_CANDIDATES=json.dumps(sigma["contacts"], ensure_ascii=False), EMD_CONTACT_CANDIDATES=json.dumps(emd["contacts"], ensure_ascii=False),
                      SIGUNGU_CD=sigma["code"], SIGUNGU_NM=sigma["name"], EMD_CD=emd["code"], EMD_NM=emd["name"],
                      LONGITUDE=None, LATITUDE=None, DISTANCE_TO_BOUNDARY_M=None, REVIEW_X_M=None, REVIEW_Y_M=None, OUTSIDE_TARGET=None,
                      SOURCE_CRS=str(config["source_coordinates"]["crs"]), SOURCE_CRS_EVIDENCE=config["source_coordinates"]["evidence"],
                      SOURCE_UNIT=config["source_coordinates"]["unit"], SOURCE_XY_ORDER=config["source_coordinates"]["xy_order"],
                      SOURCE_SCALE_X=config["source_coordinates"]["scale_x"], SOURCE_SCALE_Y=config["source_coordinates"]["scale_y"],
                      POINT_MEANING=config["source_coordinates"].get("point_meaning", "provided_point"),
                      SIGUNGU_BOUNDARY_DATE=config["boundaries"]["sigungu"]["reference_date"], EMD_BOUNDARY_DATE=config["boundaries"]["emd"]["reference_date"],
                      SIGUNGU_BOUNDARY_CRS=str(boundaries["sigungu"].crs), EMD_BOUNDARY_CRS=str(boundaries["emd"].crs),
                      SIGUNGU_CODE_SYSTEM=config["boundaries"]["sigungu"]["code_system"], EMD_CODE_SYSTEM=config["boundaries"]["emd"]["code_system"],
                      EMD_KIND=config["boundaries"]["emd"]["kind"], BOUNDARY_MODE=config["boundary_mode"], METRIC_CRS=metric_crs.to_string(),
                      MATCH_METHOD="provided_point:within;unmatched:intersects_diagnostic")
    if row.ym.startswith(INVALID) or not row.block:
        reasons.append("INVALID_LOCATION_KEY")
    if finite_float(row.x) is None or finite_float(row.y) is None:
        reasons.append("INVALID_COORDINATE")
    elif location_id not in coordinates:
        reasons.append("TRANSFORM_FAILED")
    if location_id in coordinates:
        values = coordinates[location_id]
        attributes["LONGITUDE"], attributes["LATITUDE"] = values["wgs84"]
        attributes["REVIEW_X_M"], attributes["REVIEW_Y_M"] = values["metric"]
        point = Point(*values["metric"])
        distances = []
        for name, match in (("sigungu", sigma), ("emd", emd)):
            if match["code"]:
                distances.append(point.distance(metric_geometry[name][match["code"]].boundary))
            if match["contacts"]:
                distances.append(0.0)
        attributes["DISTANCE_TO_BOUNDARY_M"] = min(distances) if distances else None
        for name, match in (("SIGUNGU", sigma), ("EMD", emd)):
            if match["count"] == 0:
                reasons.append(name + ("_BOUNDARY_CONTACT" if match["contacts"] else "_UNMATCHED"))
            elif match["count"] > 1:
                reasons.append(name + "_MULTIPLE_MATCH")
    if attributes["SIGUNGU_BOUNDARY_DATE"] != attributes["EMD_BOUNDARY_DATE"]:
        reasons.append("BOUNDARY_REFERENCE_DATE_MISMATCH")
    official = block_lookup.get((row.ym, row.block))
    official_issue = block_issues.get((row.ym, row.block))
    attributes["OFFICIAL_BLOCK_STATUS"] = "NOT_PROVIDED" if not block_supplied else "NO_APPLICABLE_RECORD"
    if official_issue:
        attributes["OFFICIAL_BLOCK_STATUS"] = official_issue
        reasons.append(official_issue)
    if official:
        official_sigma, official_emd = official
        attributes.update(OFFICIAL_BLOCK_STATUS="APPLICABLE_UNIQUE_RECORD", OFFICIAL_SIGUNGU_CD=official_sigma, OFFICIAL_EMD_CD=official_emd,
                          SIGUNGU_CD=official_sigma, EMD_CD=official_emd,
                          MATCH_METHOD="official_block_crosswalk_priority+provided_point:within;unmatched:intersects_diagnostic")
        for level, code, spatial in (("sigungu", official_sigma, sigma), ("emd", official_emd, emd)):
            names = region_names[level]
            attributes["SIGUNGU_NM" if level == "sigungu" else "EMD_NM"] = names.get(code, "")
            if code not in names:
                reasons.append("OFFICIAL_" + level.upper() + "_NOT_IN_BOUNDARY")
            elif spatial["code"] and spatial["code"] != code:
                reasons.append("OFFICIAL_SPATIAL_" + level.upper() + "_CONFLICT")
    final_sigma, final_emd = attributes["SIGUNGU_CD"], attributes["EMD_CD"]
    if not parent_supplied:
        attributes["HIERARCHY_STATUS"] = "OFFICIAL_PARENT_CROSSWALK_NOT_PROVIDED"
        reasons.append("HIERARCHY_UNVERIFIED")
    elif final_emd in parent_ambiguous:
        attributes["HIERARCHY_STATUS"] = "PARENT_CROSSWALK_NOT_UNIQUE"
        reasons.append("HIERARCHY_UNVERIFIED")
    elif not final_emd or not final_sigma:
        attributes["HIERARCHY_STATUS"] = "REGION_MATCH_HELD"
    elif final_emd not in parent_lookup:
        attributes["HIERARCHY_STATUS"] = "EMD_NOT_IN_PARENT_CROSSWALK"
        reasons.append("HIERARCHY_UNVERIFIED")
    else:
        parent = parent_lookup[final_emd]
        attributes["EMD_PARENT_SIGUNGU_CD"] = parent
        attributes["HIERARCHY_STATUS"] = "CONSISTENT" if parent == final_sigma else "MISMATCH"
        if parent != final_sigma:
            reasons.append("HIERARCHY_MISMATCH")
    if sigma["code"] and emd["code"] and emd["code"] in parent_lookup and parent_lookup[emd["code"]] != sigma["code"]:
        reasons.append("SPATIAL_HIERARCHY_MISMATCH")
    target_codes = {target["sigungu_code"] for target in config["targets"]}
    judged_codes = [code for code in (final_sigma, sigma["code"]) if code]
    if judged_codes:
        outside = any(code not in target_codes for code in judged_codes)
        attributes["OUTSIDE_TARGET"] = int(outside)
        if outside:
            reasons.append("OUTSIDE_TARGET_REGIONS")
    reasons = list(dict.fromkeys(reasons))
    attributes["MATCH_STATUS"] = "OK" if not reasons else "HOLD"
    attributes["REVIEW_REASONS"] = "|".join(reasons)
    return attributes


def map_command(config, index_requested):
    from pyproj import Transformer
    # 이 단계에서 설정 부족으로 중단되어도 index는 별도로 실행할 수 있다.
    source_crs = coordinate_settings(config)
    metric_crs = metric_settings(config, source_crs)
    if config["review"]["samples_per_category"] < 1 or config["review"]["near_boundary_m"] < 0:
        raise ValueError("review.samples_per_category는 양수, near_boundary_m는 0 이상이어야 합니다.")
    if config.get("boundary_mode") != "fixed_snapshot":
        raise ValueError("이 스크립트는 fixed_snapshot만 지원합니다. 월 당시 경계 재현은 별도 월별 경계가 필요합니다.")
    for target in config["targets"]:
        require_fields(target, ["name", "sigungu_code"], "targets")
        if not isinstance(target["sigungu_code"], str):
            raise ValueError("targets.sigungu_code는 문자열이어야 합니다.")
    index_dir, index_manifest = find_run(config, "index", index_requested)
    out = run_dir(config, "map")
    manifest = {"kind": "map", "status": "running", "normalization_version": NORMALIZATION_VERSION,
                "index_dir": str(index_dir), "input_columns": index_manifest["input_columns"]}
    dump_json(out / "manifest.json", manifest)
    print(f"출력: {out}", flush=True)
    boundaries = {name: load_boundary(config, name, out) for name in ("sigungu", "emd")}
    sigma_codes = set(boundaries["sigungu"]["region_code"])
    for target in config["targets"]:
        if target["sigungu_code"] not in sigma_codes:
            raise ValueError(f"대상 지역 코드가 시군구 경계에 없습니다: {target}")
        actual_name = boundaries["sigungu"].set_index("region_code")["region_name"][target["sigungu_code"]]
        if target["name"] not in actual_name:
            raise ValueError(f"대상 지역 이름/코드가 다릅니다: 설정={target}, 경계 이름={actual_name}")
    metric_geometry = {name: dict(zip(boundary["region_code"], boundary.to_crs(metric_crs).geometry)) for name, boundary in boundaries.items()}
    region_names = {name: dict(zip(boundary["region_code"], boundary["region_name"])) for name, boundary in boundaries.items()}
    parent_lookup, parent_ambiguous, parent_supplied = load_parent_crosswalk(config, out)
    db = readonly_db(index_dir / "location_index.sqlite")
    months = [row[0] for row in db.execute("SELECT DISTINCT ym FROM locations WHERE ym NOT LIKE '!INVALID:%' ORDER BY ym")]
    block_lookup, block_issues, block_supplied = load_block_crosswalk(config, months, out)
    transformers = {name: Transformer.from_crs(source_crs, boundary.crs, always_xy=True) for name, boundary in boundaries.items()}
    transformers["wgs84"] = Transformer.from_crs(source_crs, "EPSG:4326", always_xy=True)
    transformers["metric"] = Transformer.from_crs(source_crs, metric_crs, always_xy=True)
    map_db = sqlite3.connect(out / "mapping.sqlite")
    map_db.execute("CREATE TABLE mapping(location_id INTEGER PRIMARY KEY,ym TEXT,block TEXT,x TEXT,y TEXT,attributes TEXT,UNIQUE(ym,block,x,y))")
    fields = [*ORIGINAL, *KEYS, *MAP_FIELDS]
    main_stream, main_writer = csv_writer(out / "normalized_location_mapping.csv", fields)
    candidate_stream, candidate_writer = csv_writer(out / "spatial_match_candidates.csv", ["location_id", "layer", "predicate", "region_code", "region_name"])
    review_files = ["held_points", "unmatched_points", "multi_match_points", "boundary_contact_points", "hierarchy_mismatch_points", "outside_target_points", "invalid_coordinate_points", "official_conflict_points"]
    review_handles = {name: csv_writer(out / f"{name}.csv", fields) for name in review_files}
    sampler = ReviewSamples(config)
    counts, geo_count = Counter(), 0
    geo_stream = (out / "map_review_points.geojson").open("w", encoding="utf-8")
    geo_stream.write('{"type":"FeatureCollection","features":[')
    try:
        for frame in point_batches(db, config["chunksize"]):
            points, coordinates = prepare_points(frame, config, transformers, {name: b.crs for name, b in boundaries.items()})
            matches = {name: spatial_matches(points[name], boundary, name, candidate_writer) for name, boundary in boundaries.items()}
            database_rows = []
            for row in frame.itertuples(index=False):
                attributes = mapping_attributes(row, config, source_crs, metric_crs, boundaries, matches, coordinates,
                    parent_lookup, parent_ambiguous, parent_supplied, block_lookup, block_issues, block_supplied, metric_geometry, region_names)
                record = {**dict(zip(ORIGINAL, (row.raw_ym, row.raw_block, row.raw_x, row.raw_y))),
                          **dict(zip(KEYS, (row.ym, row.block, row.x, row.y))), **attributes}
                main_writer.writerow(record)
                reasons = attributes["REVIEW_REASONS"]
                counts[attributes["MATCH_STATUS"]] += 1
                for reason in filter(None, reasons.split("|")):
                    counts[reason] += 1
                categories = {
                    "held_points": attributes["MATCH_STATUS"] == "HOLD",
                    "unmatched_points": location_id_has_coordinates(row.location_id, coordinates) and (attributes["SIGUNGU_WITHIN_COUNT"] == 0 or attributes["EMD_WITHIN_COUNT"] == 0),
                    "multi_match_points": attributes["SIGUNGU_WITHIN_COUNT"] > 1 or attributes["EMD_WITHIN_COUNT"] > 1,
                    "boundary_contact_points": attributes["SIGUNGU_CONTACT_COUNT"] > 0 or attributes["EMD_CONTACT_COUNT"] > 0,
                    "hierarchy_mismatch_points": "HIERARCHY_" in reasons,
                    "outside_target_points": attributes["OUTSIDE_TARGET"] == 1,
                    "invalid_coordinate_points": "INVALID_" in reasons or "TRANSFORM_FAILED" in reasons,
                    "official_conflict_points": "OFFICIAL_" in reasons
                }
                for category, include in categories.items():
                    if include:
                        review_handles[category][1].writerow(record)
                sampler.consider(record)
                database_rows.append((int(row.location_id), row.ym, row.block, row.x, row.y, json.dumps(attributes, ensure_ascii=False, allow_nan=False)))
                if attributes["LONGITUDE"] is not None:
                    feature = {"type": "Feature", "geometry": {"type": "Point", "coordinates": [attributes["LONGITUDE"], attributes["LATITUDE"]]}, "properties": record}
                    if geo_count:
                        geo_stream.write(",")
                    geo_stream.write(json.dumps(feature, ensure_ascii=False, allow_nan=False))
                    geo_count += 1
            map_db.executemany("INSERT INTO mapping VALUES (?,?,?,?,?,?)", database_rows)
            map_db.commit()
            print(f"  공간 매핑 위치 수: {sum(counts[s] for s in ('OK', 'HOLD')):,}", flush=True)
        geo_stream.write("]}")
    finally:
        main_stream.close()
        candidate_stream.close()
        geo_stream.close()
        for stream, _ in review_handles.values():
            stream.close()
        map_db.close()
    # 같은 정규화 키에 속하는 원본 표기 차이를 모두 출력한다.
    db.execute("ATTACH DATABASE ? AS mapped", ((out / "mapping.sqlite").resolve().as_uri() + "?mode=ro",))
    alias_stream, alias_writer = csv_writer(out / "admin_mapping.csv", fields)
    with alias_stream:
        query = """SELECT DISTINCT o.raw_ym,o.raw_block,o.raw_x,o.raw_y,o.ym,o.block,o.x,o.y,m.attributes
            FROM observations o JOIN mapped.mapping m ON m.ym=o.ym AND m.block=o.block AND m.x=o.x AND m.y=o.y
            ORDER BY o.ym,o.block,o.x,o.y,o.raw_ym,o.raw_block,o.raw_x,o.raw_y"""
        for raw_ym, raw_block, raw_x, raw_y, ym, block, x, y, attrs in db.execute(query):
            alias_writer.writerow({**dict(zip(ORIGINAL, (raw_ym, raw_block, raw_x, raw_y))), **dict(zip(KEYS, (ym, block, x, y))), **json.loads(attrs)})
    expected_locations = db.execute("SELECT COUNT(*) FROM locations").fetchone()[0]
    db.close()
    if sum(counts[s] for s in ("OK", "HOLD")) != expected_locations:
        raise ValueError("행정구역 대응표의 행 수가 위치 인덱스의 행 수와 다릅니다.")
    sampler.save(out)
    manifest.update(status="complete", locations=sum(counts[s] for s in ("OK", "HOLD")), geojson_points=geo_count,
                    diagnostics=dict(counts), boundary_mode="fixed_snapshot", point_semantics="제공된 점이 속한 행정구역")
    dump_json(out / "manifest.json", manifest)
    print(f"공간 결합 완료: {out}\nOK 후보와 HOLD 검토 목록을 함께 저장했습니다.")
    return out


def location_id_has_coordinates(location_id, coordinates):
    return int(location_id) in coordinates


def exact_add(left, right):
    """인구 합계에 Decimal 기본 정밀도에 따른 반올림이 생기지 않도록 한다."""
    ltuple, rtuple = left.as_tuple(), right.as_tuple()
    exponent = min(ltuple.exponent, rtuple.exponent)
    precision = max(len(ltuple.digits) + ltuple.exponent - exponent, len(rtuple.digits) + rtuple.exponent - exponent) + 2
    with localcontext() as context:
        context.prec = max(28, precision)
        return left + right


def numeric_profile(values):
    total, numeric, empty, nonnumeric = Decimal(0), 0, 0, 0
    for value in values:
        text = str(value).strip()
        if not text:
            empty += 1
            continue
        try:
            number = Decimal(text)
            if not number.is_finite() or len(number.as_tuple().digits) > 2048 or abs(number.as_tuple().exponent) > 10000:
                raise InvalidOperation
        except InvalidOperation:
            nonnumeric += 1
            continue
        total = exact_add(total, number)
        numeric += 1
    return total, numeric, empty, nonnumeric


def merge_chunk_with_mapping(frame, columns, mapping_db):
    """해당 청크의 키만 SQLite에서 조회하고 many_to_one으로 결합한다."""
    join_columns = ["_SK_JOIN_YM", "_SK_JOIN_BLOCK", "_SK_JOIN_X", "_SK_JOIN_Y"]
    reserved = {*join_columns, "_SK_SOURCE_POS", "_merge"}
    reserved |= {"ADM_" + field for field in MAP_FIELDS}
    conflicts = reserved.intersection(frame.columns)
    if conflicts:
        raise ValueError(f"결합에 사용하는 예약 열 이름이 원본 열과 겹칩니다: {sorted(conflicts)}")
    original = frame.reset_index(drop=True).copy()
    left = original.copy()
    normalized = [normalize_values(values) for values in original[columns].itertuples(index=False, name=None)]
    for offset, column in enumerate(join_columns):
        left[column] = [values[offset] for values in normalized]
    left["_SK_SOURCE_POS"] = list(range(len(left)))
    mapping_db.execute("CREATE TEMP TABLE IF NOT EXISTS requested_keys(ym TEXT,block TEXT,x TEXT,y TEXT,PRIMARY KEY(ym,block,x,y))")
    mapping_db.execute("DELETE FROM requested_keys")
    mapping_db.executemany("INSERT OR IGNORE INTO requested_keys VALUES (?,?,?,?)", normalized)
    rows = []
    for ym, block, x, y, attrs in mapping_db.execute("""SELECT m.ym,m.block,m.x,m.y,m.attributes FROM mapping m
        JOIN requested_keys r ON m.ym=r.ym AND m.block=r.block AND m.x=r.x AND m.y=r.y"""):
        rows.append({**dict(zip(join_columns, (ym, block, x, y))), **{"ADM_" + k: v for k, v in json.loads(attrs).items()}})
    right = pd.DataFrame(rows, columns=[*join_columns, *("ADM_" + field for field in MAP_FIELDS)])
    result = left.merge(right, on=join_columns, how="left", validate="many_to_one", indicator=True, sort=False)
    if len(result) != len(original) or result["_SK_SOURCE_POS"].tolist() != list(range(len(original))):
        raise ValueError("결합 전후 행 수 또는 원본 행 순서가 달라졌습니다.")
    if not result[original.columns].reset_index(drop=True).equals(original):
        raise ValueError("결합 전후 원본 값이 달라졌습니다.")
    not_found = result["_merge"].eq("left_only")
    result.loc[not_found, "ADM_MATCH_STATUS"] = "NOT_IN_INDEX"
    result.loc[not_found, "ADM_REVIEW_REASONS"] = "NORMALIZED_LOCATION_KEY_NOT_IN_MAPPING"
    result = result.drop(columns=[*join_columns, "_SK_SOURCE_POS", "_merge"])
    return result, not_found.to_numpy()


def attach_command(config, source_path, member, map_requested, population_requested):
    map_dir, manifest = find_run(config, "map", map_requested)
    if config["columns"] != manifest["input_columns"]:
        raise ValueError("현재 입력 열 설정이 index/map에 사용한 설정과 다릅니다.")
    source = {"path": str(Path(source_path).resolve()), "member": member or ""}
    table, _ = table_for(Path(member or source_path).name, config)
    population = population_requested or (table.get("population_columns", []) if table else [])
    if not population:
        raise ValueError("인구 값 보존을 확인할 실제 열 이름을 tables.population_columns 또는 --population-columns로 지정해주세요. 열 이름을 추측하지 않습니다.")
    requested = [config["columns"][role] for role in ROLES]
    encoding, sep, header = detect_csv(source, config, requested + population)
    columns = resolve_columns(header, requested)
    population_columns = resolve_columns(header, population)
    out = run_dir(config, "attach")
    dump_json(out / "manifest.json", {"kind": "attach", "status": "running", "normalization_version": NORMALIZATION_VERSION})
    mapping_db = readonly_db(map_dir / "mapping.sqlite")
    totals = {column: [Decimal(0), 0, 0, 0] for column in population_columns}
    n_before, n_after, n_not_found, n_hold, first = 0, 0, 0, 0, True
    output_path = out / "source_with_admin.csv"
    missing_path = out / "not_in_index_rows.csv"
    try:
        with open_binary(source) as stream, output_path.open("x", encoding="utf-8-sig", newline="") as output, missing_path.open("x", encoding="utf-8-sig", newline="") as missing_output:
            reader = pd.read_csv(stream, encoding=encoding, sep=sep, dtype=str, keep_default_na=False, na_filter=False,
                                 chunksize=config["chunksize"], on_bad_lines="error")
            for frame in reader:
                result, missing = merge_chunk_with_mapping(frame, columns, mapping_db)
                for column in population_columns:
                    before, after = numeric_profile(frame[column]), numeric_profile(result[column])
                    if before != after:
                        raise ValueError(f"인구 값 합계/결측/비수치 개수가 결합 후 달라졌습니다: {column}")
                    totals[column][0] = exact_add(totals[column][0], before[0])
                    for offset in (1, 2, 3):
                        totals[column][offset] += before[offset]
                result.to_csv(output, index=False, header=first)
                result.loc[missing].to_csv(missing_output, index=False, header=first)
                first = False
                n_before += len(frame)
                n_after += len(result)
                n_not_found += int(missing.sum())
                n_hold += int(result["ADM_MATCH_STATUS"].eq("HOLD").sum())
                print(f"  행정구역 결합 행 수: {n_after:,}", flush=True)
            if first:
                # 데이터 행이 없는 파일에도 열 이름을 남긴다.
                pd.DataFrame(columns=[*header, *("ADM_" + field for field in MAP_FIELDS)]).to_csv(output, index=False)
                pd.DataFrame(columns=[*header, *("ADM_" + field for field in MAP_FIELDS)]).to_csv(missing_output, index=False)
        if n_before != n_after:
            raise ValueError("결합 전후 전체 행 수가 달라졌습니다.")
        report = {"source": source, "mapping_dir": str(map_dir), "input_encoding": encoding, "input_separator": repr(sep),
                  "rows_before": n_before, "rows_after": n_after, "rows_preserved": True, "all_original_cells_preserved": True,
                  "held_rows_preserved": n_hold, "not_in_index_rows_preserved": n_not_found, "population_fields": {},
                  "note": "인구 열의 문자열 값은 모든 행에서 그대로 보존됨. 아래 합계는 유한 숫자 값의 합이며 비수치/빈 값은 별도 개수로 기록됨."}
        for column, values in totals.items():
            report["population_fields"][column] = {"sum_before": str(values[0]), "sum_after": str(values[0]), "numeric_rows": values[1], "empty_rows": values[2], "nonnumeric_rows": values[3], "preserved": True}
        dump_json(out / "attach_validation.json", report)
        dump_json(out / "manifest.json", {"kind": "attach", "status": "complete", "normalization_version": NORMALIZATION_VERSION, "rows": n_after, "map_dir": str(map_dir)})
    finally:
        mapping_db.close()
    print(f"원본에 행정구역 붙이기 완료: {output_path}\n원본 파일을 덮어쓰지 않았습니다.")
    return out


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=["index", "map", "attach"])
    parser.add_argument("--config", type=Path, default=None, help="별도 설정 JSON. 생략하면 공통 프로젝트 설정 사용")
    parser.add_argument("--index-dir", default="latest", help="map에 사용할 완료된 index 폴더 또는 latest")
    parser.add_argument("--map-dir", default="latest", help="attach에 사용할 완료된 map 폴더 또는 latest")
    parser.add_argument("--source", help="attach 대상 원본 CSV 또는 ZIP 파일의 절대 경로")
    parser.add_argument("--member", help="ZIP 내부 CSV의 경로; 압축을 풀지 않고 읽음")
    parser.add_argument("--population-columns", nargs="+", help="보존 확인 대상 인구 열의 실제 이름")
    args = parser.parse_args()
    if args.config:
        config = json.loads(args.config.read_text(encoding="utf-8-sig"))
    else:
        from project_settings import build_config
        config = build_config()
    if isinstance(config.get("chunksize"), bool) or not isinstance(config.get("chunksize"), int) or config["chunksize"] < 1:
        raise ValueError("chunksize는 양의 정수여야 합니다.")
    if set(config["columns"]) != set(ROLES) or len(set(config["columns"].values())) != len(ROLES):
        raise ValueError("columns에는 month/block/x/y에 해당하는 서로 다른 네 열 이름이 필요합니다.")
    if args.command == "index":
        index_command(config)
    elif args.command == "map":
        map_command(config, args.index_dir)
    else:
        if not args.source:
            parser.error("attach에는 --source가 필요합니다.")
        attach_command(config, args.source, args.member, args.map_dir, args.population_columns)


if __name__ == "__main__":
    try:
        main()
    except (ValueError, FileNotFoundError, KeyError, sqlite3.Error) as error:
        print(f"중단: {error}", file=sys.stderr)
        sys.exit(2)
